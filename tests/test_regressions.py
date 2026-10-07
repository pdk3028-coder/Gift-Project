"""Regression tests use a temporary DB and never touch employees.db."""
import io
import os
from pathlib import Path
import sqlite3
import tempfile
import unittest
from unittest.mock import patch

import html5lib
import pandas as pd
from flask import Flask
from flask.sessions import SecureCookieSessionInterface
from openpyxl import load_workbook
from PIL import Image

_SANDBOX = tempfile.TemporaryDirectory(prefix='gift_tests_')
os.environ['GIFT_DATABASE'] = str(Path(_SANDBOX.name) / 'employees.db')
os.environ['GIFT_INSTANCE'] = str(Path(_SANDBOX.name) / 'instance')
os.environ['GIFT_SECRET_KEY'] = 'test-only-secret-' + 'a' * 48
os.environ['GIFT_TRUSTED_PROXY_HOPS'] = '0'
os.environ['GIFT_RATE_LIMIT_STORAGE'] = 'memory://'
import database
import app as module

PASSWORD = 'test-admin-password-2026'


class RegressionTests(unittest.TestCase):
    def setUp(self):
        self.folder = tempfile.TemporaryDirectory(dir=_SANDBOX.name)
        self.addCleanup(self.folder.cleanup)
        self.root = Path(self.folder.name)
        database.DB_NAME = str(self.root / 'employees.db')
        database.init_db()
        database.set_admin_password(PASSWORD)
        self.app = module.app
        images = self.root / 'uploads' / 'gift_images'
        images.mkdir(parents=True)
        self.app.config.update(TESTING=True, UPLOAD_FOLDER=str(images.parent), GIFT_IMAGE_FOLDER=str(images),
                               MAX_CONTENT_LENGTH=10*1024*1024)
        module.limiter.reset()
        with database.transaction() as conn:
            conn.execute('''INSERT INTO employees(name,emp_id,phone,address_main,address_main_detail,zipcode,
                           gift_address,gift_zipcode,gift_receiver) VALUES(?,?,?,?,?,?,?,?,?)''',
                         ('Test', '00123', '010-1234-5678', 'Original', '101', '01234', 'Gift address', '01234', 'Test'))
        self.client = self.app.test_client()

    def token(self, client=None):
        client = client or self.client
        client.get('/admin/login')
        with client.session_transaction() as session:
            return session['_csrf']

    def post(self, path, data=None, client=None, **kwargs):
        client = client or self.client
        fields = dict(data or {})
        fields['csrf_token'] = self.token(client)
        return client.post(path, data=fields, **kwargs)

    def admin(self, client=None):
        client = client or self.client
        response = self.post('/admin/login', {'password': PASSWORD}, client=client)
        self.assertTrue(response.location.endswith('/admin'))
        return client

    def employee(self):
        response = self.post('/login', {'emp_id': '00123', 'agree_privacy': 'on'})
        self.assertTrue(response.location.endswith('/dashboard'))

    def test_old_forged_cookie_rejected(self):
        fake = Flask('fake'); fake.secret_key = 'your_secret_key_here'
        token = SecureCookieSessionInterface().get_signing_serializer(fake).dumps({'is_admin': True})
        self.client.set_cookie('session', token)
        self.assertEqual(self.client.get('/admin').status_code, 302)

    def test_path_escape_and_active_content_rejected(self):
        (self.root / 'uploads' / 'marker.txt').write_text('secret')
        for filename in ['..%5cmarker.txt', '..%2fmarker.txt', 'test.html', 'test.svg']:
            self.assertEqual(self.client.get('/uploads/gift_images/' + filename).status_code, 404)

    def test_csrf_required_on_every_post_route(self):
        self.admin()
        for path in ['/admin/settings','/admin/upload','/admin/reset','/admin/gifts/add',
                     '/admin/gifts/delete/1','/info_update','/gift_update','/login','/logout','/admin/logout']:
            with self.subTest(path=path):
                self.assertEqual(self.client.post(path).status_code, 400)
        self.assertEqual(self.client.post('/admin/settings', data={'csrf_token':'한글'}).status_code, 400)
        self.assertEqual(self.client.get('/logout').status_code, 405)

    def test_employee_id_login_retained_and_missing_password_safe(self):
        self.employee()
        self.assertEqual(self.client.get('/info_update').status_code, 200)
        with self.client.session_transaction() as session:
            self.assertTrue(session.permanent)
        self.assertEqual(self.post('/admin/login').status_code, 302)
        self.assertFalse(database.verify_admin_password(None))

    def test_forwarded_header_cannot_bypass_limit(self):
        for number in range(10):
            self.assertEqual(self.post('/admin/login', {'password':'wrong'},
                             headers={'X-Forwarded-For': f'192.0.2.{number}'}).status_code, 302)
        self.assertEqual(self.post('/admin/login', {'password':'wrong'},
                         headers={'X-Forwarded-For':'198.51.100.1'}).status_code, 429)

    def test_empty_and_no_change_preserve_fields(self):
        self.employee()
        self.post('/info_update')
        self.assertEqual(database.get_employee_by_id('00123')['phone'], '010-1234-5678')
        self.post('/info_update', {'no_change':'on', 'phone':'changed'})
        self.assertEqual(database.get_employee_by_id('00123')['phone'], '010-1234-5678')
        self.post('/gift_update', {'same_address':'on','gift_address':'changed'})
        self.assertEqual(database.get_employee_by_id('00123')['gift_address'], 'Gift address')

    def test_valid_information_update(self):
        self.employee()
        self.post('/info_update', {'phone':'010-9876-5432','address_main':'New address',
                                 'address_main_detail':'202','zipcode':'02345'})
        self.assertEqual(database.get_employee_by_id('00123')['address_main'], 'New address')

    def test_gift_validation_and_history(self):
        database.add_gift_option('Gift A','','')
        gift = database.get_gift_options()[0]
        self.employee()
        self.post('/gift_update', {'same_address':'on','selected_gift_id':'999999'})
        self.assertIsNone(database.get_employee_by_id('00123')['selected_gift_id'])
        self.post('/gift_update', {'same_address':'on','selected_gift_id':str(gift['id'])})
        self.assertEqual(database.get_employee_by_id('00123')['selected_gift_id'], gift['id'])
        database.delete_gift_option(gift['id'])
        self.assertEqual(database.get_gift_options(), [])
        self.assertEqual(database.get_all_employees().iloc[0]['gift_name'], 'Gift A')
        with self.assertRaises(sqlite3.IntegrityError):
            database.update_employee_info('00123', {'selected_gift_id':gift['id']})
        # A failed transaction must not leave the DB locked.
        database.update_employee_info('00123', {'phone':'01011112222'})

    def test_export_literal_text_and_no_disk_file(self):
        database.update_employee_info('00123', {'phone':'=1+1','address_main':'+CMD','gift_receiver':'@SUM(A1)'})
        self.admin()
        response = self.client.get('/admin/download')
        with io.BytesIO(response.data) as buffer:
            workbook = load_workbook(buffer)
            sheet = workbook.active
            headers = {cell.value:cell.column for cell in sheet[1]}
            for field in ['phone','address_main','gift_receiver']:
                self.assertEqual(sheet.cell(2,headers[field]).data_type,'s')
            self.assertEqual(sheet.cell(2,headers['emp_id']).value,'00123')
            workbook.close()
        response.close()
        self.assertEqual(list((self.root/'uploads').glob('*.xlsx')), [])

    def test_image_validation_reencoding_and_size(self):
        self.admin()
        self.post('/admin/gifts/add', {'name':'Bad','description':'','image':(io.BytesIO(b'<html>bad</html>'),'x.jpg')})
        self.assertEqual(database.get_gift_options(), [])
        output=io.BytesIO();Image.new('RGB',(4,4),'red').save(output,format='PNG');output.seek(0)
        self.post('/admin/gifts/add', {'name':'Good','description':'','image':(output,'한글.png')})
        gift=database.get_gift_options()[0]
        response=self.client.get('/'+gift['image_path'])
        self.assertEqual(response.mimetype,'image/jpeg')
        self.assertEqual(response.headers['X-Content-Type-Options'],'nosniff')
        response.close()
        self.app.config['MAX_CONTENT_LENGTH']=100
        self.assertEqual(self.client.post('/admin/upload',data=b'x'*101).status_code,413)

    def roster(self, **overrides):
        row={'사번':'00123','성명':'Test','핸드폰':'01011112222','현거주지':'New','우편번호-현':'01234'}
        row.update(overrides)
        return row

    def import_rows(self, rows, name='input.xlsx'):
        path=self.root/name
        pd.DataFrame(rows).to_excel(path,index=False)
        return database.upsert_employees_from_excel(path)

    def test_import_schema_duplicate_and_atomicity(self):
        for rows in [[{'unexpected':'value'}], [self.roster(),self.roster()],
                     [self.roster(),self.roster(사번='')], [self.roster(**{'우편번호-현':'bad'})]]:
            with self.subTest(rows=len(rows)):
                with self.assertRaises(ValueError):self.import_rows(rows)
                self.assertEqual(database.get_employee_by_id('00123')['address_main'],'Original')
                self.assertIsNone(database.get_employee_by_id(''))
        self.assertEqual(self.import_rows([self.roster()]),1)
        self.assertEqual(database.get_employee_by_id('00123')['address_main'],'New')
        self.import_rows([self.roster(핸드폰='',현거주지='',**{'우편번호-현':''})])
        self.assertEqual(database.get_employee_by_id('00123')['phone'],'01011112222')

    def test_upload_rejects_bad_type_and_discards_temporary_source(self):
        self.admin()
        before=list((self.root/'uploads').iterdir())
        self.post('/admin/upload',{'file':(io.BytesIO(b'bad'),'file.html')})
        buffer=io.BytesIO();pd.DataFrame([self.roster()]).to_excel(buffer,index=False);buffer.seek(0)
        self.post('/admin/upload',{'file':(buffer,'사원명부.xlsx')})
        self.assertEqual(database.get_employee_by_id('00123')['address_main'],'New')
        self.assertEqual(list((self.root/'uploads').iterdir()),before)

    def test_import_registered_address_fallback_keeps_postcode_paired(self):
        fallback = {'주민등록주소지': 'Registered', '우편번호-주': '04567'}
        self.import_rows([self.roster(**fallback)])
        employee = database.get_employee_by_id('00123')
        self.assertEqual((employee['address_main'], employee['zipcode']), ('New', '01234'))
        self.import_rows([self.roster(현거주지='  ', **fallback)])
        employee = database.get_employee_by_id('00123')
        self.assertEqual((employee['address_main'], employee['zipcode']), ('Registered', '04567'))
        fallback['우편번호-주'] = ''
        self.import_rows([self.roster(현거주지='', **fallback)])
        employee = database.get_employee_by_id('00123')
        self.assertEqual((employee['address_main'], employee['zipcode']), ('Registered', ''))

    def test_import_invalid_fallback_is_atomic(self):
        for extra in ({'주민등록주소지': 'Registered'},
                      {'주민등록주소지': 'Registered', '우편번호-주': 'bad'}):
            with self.subTest(extra=extra):
                with self.assertRaises(ValueError):
                    self.import_rows([self.roster(현거주지='', **extra)])
                self.assertEqual(database.get_employee_by_id('00123')['address_main'], 'Original')

    def test_duplicate_headers_rejected(self):
        path=self.root/'duplicates.xlsx'
        row=self.roster()
        frame=pd.DataFrame([list(row.values())+['other']],columns=list(row)+['사번'])
        frame.to_excel(path,index=False)
        with self.assertRaises(ValueError):database.upsert_employees_from_excel(path)
        self.assertEqual(database.get_employee_by_id('00123')['address_main'],'Original')

    def test_import_preserves_current_and_legacy_postcodes(self):
        for source, expected in [('01234','01234'),('689853','689853'),('689-853','689853')]:
            with self.subTest(source=source):
                self.import_rows([self.roster(**{'우편번호-현':source})])
                self.assertEqual(database.get_employee_by_id('00123')['zipcode'],expected)
        for invalid in ['1234','1234567','12A45','12-345','１２３４５']:
            with self.subTest(invalid=invalid):
                with self.assertRaises(ValueError):
                    self.import_rows([self.roster(**{'우편번호-현':invalid})])
                self.assertEqual(database.get_employee_by_id('00123')['zipcode'],'689853')

    def test_import_into_legacy_database_with_required_ssn(self):
        with database.transaction() as conn:
            schema=conn.execute("SELECT sql FROM sqlite_master WHERE name='employees'").fetchone()[0]
        schema=schema.replace("ssn TEXT DEFAULT ''", 'ssn TEXT NOT NULL')
        database.DB_NAME=str(self.root/'legacy.db')
        with database.transaction() as conn:
            conn.execute(schema)
        database.init_db()
        self.assertEqual(self.import_rows([self.roster()]),1)
        self.assertEqual(database.get_employee_by_id('00123')['ssn'],'')
        # SQLite checks NOT NULL before resolving the UPSERT conflict, so cover
        # re-importing existing staff as well as inserting new staff.
        self.assertEqual(self.import_rows([self.roster(현거주지='Updated')]),1)
        employee=database.get_employee_by_id('00123')
        self.assertEqual(employee['address_main'],'Updated')
        self.assertEqual(employee['ssn'],'')

    def test_reset_preserves_password_and_no_default_created(self):
        self.admin()
        self.post('/admin/reset',{'password':PASSWORD})
        self.assertTrue(database.verify_admin_password(PASSWORD))
        database.init_db()
        self.assertTrue(database.verify_admin_password(PASSWORD))
        self.assertFalse(database.verify_admin_password('admin1234'))
        self.assertIsNone(database.get_employee_by_id('00123'))
        database.DB_NAME=str(self.root/'fresh.db');database.init_db()
        self.assertEqual(database.get_setting('admin_password',''),'')

    def test_password_change_revokes_other_admin_session(self):
        other=self.app.test_client();self.admin(other);self.admin()
        self.post('/admin/settings',{'new_password':'changed-password-2026','enable_info':'on'})
        self.assertEqual(other.get('/admin').status_code,302)
        self.assertEqual(self.client.get('/admin').status_code,200)

    def test_templates_are_well_formed_with_csrf(self):
        self.employee()
        database.add_gift_option('Gift','','')
        for path in ['/dashboard','/info_update','/gift_update']:
            self.assert_valid_html(path)
        self.admin()
        for path in ['/admin','/admin/gifts','/admin/login']:
            self.assert_valid_html(path)
        client=self.app.test_client()
        self.assert_valid_html('/',client)
        self.assert_valid_html('/admin/login',client)

    def assert_valid_html(self,path,client=None):
        response=(client or self.client).get(path)
        self.assertEqual(response.status_code,200)
        parser=html5lib.HTMLParser(namespaceHTMLElements=False)
        document=parser.parse(response.data.decode())
        self.assertEqual(parser.errors,[],path)
        for form in document.iter('form'):
            if form.get('method','').lower()=='post':
                self.assertTrue(any(element.get('name')=='csrf_token' for element in form.iter()),path)
                self.assertTrue(any(element.tag=='button' and element.get('type')=='submit' for element in form.iter()),path)

    def test_cleanup_dry_run_and_reference_preservation(self):
        root=Path(self.app.config['UPLOAD_FOLDER']);images=Path(self.app.config['GIFT_IMAGE_FOLDER'])
        old=root/'old.xlsx';old.write_bytes(b'legacy')
        orphan=images/'orphan.jpg';orphan.write_bytes(b'orphan')
        keep=images/'keep.jpg';keep.write_bytes(b'keep')
        database.add_gift_option('Keep','','uploads/gift_images/keep.jpg')
        runner=self.app.test_cli_runner()
        self.assertEqual(runner.invoke(args=['cleanup-uploads']).exit_code,0)
        self.assertTrue(old.exists())
        self.assertEqual(runner.invoke(args=['cleanup-uploads','--apply']).exit_code,0)
        self.assertFalse(old.exists());self.assertFalse(orphan.exists());self.assertTrue(keep.exists())

    def test_absolute_db_and_secret_persistence(self):
        self.assertTrue(Path(database.DB_NAME).is_absolute())
        original=Path.cwd()
        try:
            os.chdir(self.root)
            self.assertIsNotNone(database.get_employee_by_id('00123'))
        finally:os.chdir(original)
        with patch.dict(os.environ):
            os.environ.pop('GIFT_SECRET_KEY',None)
            self.assertEqual(module.load_secret(),module.load_secret())
            os.environ['GIFT_SECRET_KEY']='your_secret_key_here'
            with self.assertRaises(RuntimeError):module.load_secret()


if __name__=='__main__':
    try:
        unittest.main()
    finally:
        _SANDBOX.cleanup()
